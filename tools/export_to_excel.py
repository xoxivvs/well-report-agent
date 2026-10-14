"""
Инструмент: экспорт данных отчёта в Excel.
Версия 0.2: поддержка Volve (well_name вместо well_id).
"""

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if __package__ in (None, ""):  # прямой запуск: python tools/export_to_excel.py
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.config import EXCEL_DIR

HEADER_FILL = PatternFill("solid", fgColor="2C3E50")
HEADER_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D6E4F0")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="1F3864")
BODY_FONT = Font(name="Calibri", size=11)
BORDER = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def _style_body(ws, start_row, end_row, ncols):
    for r in range(start_row, end_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = LEFT


def _autosize(ws, min_width=12, max_width=45):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = min_width
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, min(len(str(cell.value)), max_width))
        ws.column_dimensions[col_letter].width = max_len + 2


def _write_section(ws, row, title):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cell = ws.cell(row=row, column=1, value=title)
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    cell.alignment = LEFT
    return row + 1


def export_to_excel(well_identifier, period, production, telemetry, geo):
    """
    well_identifier: int (синтетика) или str (Volve — например '15/9-F-14').
    """
    try:
        os.makedirs(EXCEL_DIR, exist_ok=True)

        # Определяем источник
        is_volve = (
            production.get("status") == "success"
            and "well_name" in (production.get("data") or {})
        )
        if is_volve:
            well_label = production["data"]["well_name"]
        else:
            well_label = f"№{well_identifier}"

        safe_name = str(well_identifier).replace("/", "_").replace(" ", "_").replace("№", "")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = os.path.join(EXCEL_DIR, f"report_well{safe_name}_{period}_{timestamp}.xlsx")

        unit = "м³/сут" if is_volve else "т/сут"
        volume_unit = "м³" if is_volve else "т"

        wb = Workbook()

        # ===== Лист 1: Сводка =====
        ws = wb.active
        ws.title = "Сводка"

        ws.merge_cells("A1:B1")
        ws["A1"] = f"Отчёт по скважине {well_label}"
        ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F3864")
        ws["A1"].alignment = CENTER
        ws.row_dimensions[1].height = 28

        ws["A2"] = "Период:"
        ws["B2"] = period
        ws["A3"] = "Дата формирования:"
        ws["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        if is_volve:
            ws["A4"] = "Источник данных:"
            ws["B4"] = "Equinor Volve Open Dataset (реальные данные)"
        for r in (2, 3, 4):
            ws.cell(row=r, column=1).font = Font(bold=True)
            ws.cell(row=r, column=2).font = BODY_FONT

        row = 6

        # Добыча
        if production["status"] == "success":
            agg = production["data"]["aggregates"]
            trend = production["data"]["trend"]

            row = _write_section(ws, row, "ПОКАЗАТЕЛИ ДОБЫЧИ")
            _style_header(ws, row, 2)
            ws.cell(row=row, column=1, value="Показатель").font = HEADER_FONT
            ws.cell(row=row, column=2, value="Значение").font = HEADER_FONT
            row += 1

            avg_oil = agg.get("avg_oil") if is_volve else agg.get("avg_oil_rate")
            avg_wc = agg.get("avg_water_cut", 0)
            days = agg.get("days_count")
            working = agg.get("working_days", days)
            delta = trend.get("delta_oil") if is_volve else trend.get("delta_oil_rate")

            rows_data = [
                ("Средний дебит нефти", f"{avg_oil} {unit}"),
                ("Средняя обводнённость", f"{avg_wc} %"),
                ("Суммарная добыча нефти", f"{agg['total_oil']} {volume_unit}"),
                ("Дней в отчёте", f"{days} (рабочих: {working})"),
                ("Тренд дебита", f"{trend['direction']} ({delta})"),
            ]
            start = row
            for name, value in rows_data:
                ws.cell(row=row, column=1, value=name)
                ws.cell(row=row, column=2, value=value)
                row += 1
            _style_body(ws, start, row - 1, 2)
            row += 1

        # УЭЦН
        if telemetry.get("status") == "success":
            t = telemetry["data"]
            row = _write_section(ws, row, "СОСТОЯНИЕ ОБОРУДОВАНИЯ (УЭЦН — ДЕМО)")
            _style_header(ws, row, 2)
            ws.cell(row=row, column=1, value="Показатель").font = HEADER_FONT
            ws.cell(row=row, column=2, value="Значение").font = HEADER_FONT
            row += 1

            alarms_str = ", ".join(f"{k}: {v}" for k, v in t["alarms_summary"].items()) or "—"
            rows_data = [
                ("Наработка на конец периода", f"{t['runtime_end']} ч"),
                ("Тревог всего", str(len(t["alarms"]))),
                ("Типы тревог", alarms_str),
            ]
            start = row
            for name, value in rows_data:
                ws.cell(row=row, column=1, value=name)
                ws.cell(row=row, column=2, value=value)
                row += 1
            _style_body(ws, start, row - 1, 2)
            row += 1

        # Геология
        if geo.get("status") == "success":
            d = geo["data"]
            row = _write_section(ws, row, "ГЕОЛОГИЧЕСКИЕ ПАРАМЕТРЫ")
            _style_header(ws, row, 2)
            ws.cell(row=row, column=1, value="Показатель").font = HEADER_FONT
            ws.cell(row=row, column=2, value="Значение").font = HEADER_FONT
            row += 1

            rows_data = [
                ("Пластовое давление", f"{d['reservoir_pressure']} атм"),
                ("Режим работы", d["mode"]),
                ("Дата отчёта", d["date"] or "—"),
            ]
            start = row
            for name, value in rows_data:
                ws.cell(row=row, column=1, value=name)
                ws.cell(row=row, column=2, value=value)
                row += 1
            _style_body(ws, start, row - 1, 2)

        _autosize(ws)

        # ===== Лист 2: Добыча =====
        if production["status"] == "success":
            ws2 = wb.create_sheet("Добыча")

            if is_volve:
                headers = ["Дата", "Часы работы", "Дебит нефти, м³/сут",
                           "Дебит газа, м³/сут", "Дебит воды, м³/сут",
                           "Обводнённость, %", "Забойное давление, атм",
                           "Устьевое давление, атм"]
                ws2.append(headers)
                _style_header(ws2, 1, len(headers))

                for r in production["data"]["records"]:
                    ws2.append([
                        r.get("date"),
                        r.get("on_stream_hours"),
                        r.get("oil_volume"),
                        r.get("gas_volume"),
                        r.get("water_volume"),
                        r.get("water_cut"),
                        r.get("downhole_pressure"),
                        r.get("wellhead_pressure"),
                    ])
                _style_body(ws2, 2, ws2.max_row, len(headers))
            else:
                headers = ["Дата", "Дебит жидкости, т/сут", "Дебит нефти, т/сут", "Обводнённость, %"]
                ws2.append(headers)
                _style_header(ws2, 1, len(headers))
                for r in production["data"]["records"]:
                    ws2.append([r["date"], r["liquid_rate"], r["oil_rate"], r["water_cut"]])
                _style_body(ws2, 2, ws2.max_row, len(headers))

            _autosize(ws2)

        # ===== Лист 3: Телеметрия =====
        if telemetry.get("status") == "success":
            ws3 = wb.create_sheet("Телеметрия")
            alarms = telemetry["data"]["alarms"]
            if alarms:
                headers = ["Время", "Наработка, ч", "Код тревоги"]
                ws3.append(headers)
                _style_header(ws3, 1, len(headers))
                for a in alarms:
                    ws3.append([a["timestamp"], a["runtime_hours"], a["alarm_code"]])
                _style_body(ws3, 2, ws3.max_row, len(headers))
            else:
                ws3["A1"] = "Тревог за период не зафиксировано"
                ws3["A1"].font = Font(italic=True, color="777777")
            _autosize(ws3)

        # ===== Лист 4: Геология =====
        if geo.get("status") == "success":
            ws4 = wb.create_sheet("Геология")
            d = geo["data"]
            ws4.append(["Показатель", "Значение"])
            _style_header(ws4, 1, 2)
            for name, value in [
                ("Скважина", well_label),
                ("Период", period),
                ("Дата отчёта", d.get("date")),
                ("Пластовое давление, атм", d.get("reservoir_pressure")),
                ("Режим", d.get("mode")),
            ]:
                ws4.append([name, value])
            _style_body(ws4, 2, ws4.max_row, 2)
            _autosize(ws4)

        wb.save(path)

        return {
            "status": "success",
            "data": {"xlsx_path": path},
            "error": None,
        }

    except Exception as e:
        import traceback
        return {
            "status": "error",
            "data": None,
            "error": f"Ошибка экспорта в Excel: {e}\n{traceback.format_exc()}",
        }


if __name__ == "__main__":
    import sys
    sys.path.insert(0, ".")
    from tools.get_geo_report import get_geo_report
    from tools.get_production_data import get_production_data
    from tools.get_telemetry import get_telemetry

    print("=== Volve ===")
    well_name, period = "15/9-F-14", "2014-06"
    prod = get_production_data(well_name, period, source="volve")
    tel = get_telemetry(5, period)
    geo = get_geo_report(5, period)
    result = export_to_excel(well_name, period, prod, tel, geo)
    print(result)

    if result["status"] == "success":
        print(f"\n📊 Excel: {result['data']['xlsx_path']}")

    print("\n=== Synthetic ===")
    prod_syn = get_production_data(5, "2026-09", source="synthetic")
    result_syn = export_to_excel(5, "2026-09", prod_syn, tel, geo)
    print(f"Статус: {result_syn['status']}")
    if result_syn["status"] == "success":
        print(f"📊 Excel: {result_syn['data']['xlsx_path']}")
